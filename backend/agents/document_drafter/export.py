"""
=============================================================================
LEXIS Document Drafter — PDF & DOCX Export Service (Module 6)
=============================================================================

Generates downloadable PDF and DOCX binary files from assembled document text.

Uses:
  - reportlab for PDF generation (pure Python, zero external binary dependencies)
  - python-docx for Microsoft Word (.docx) generation
"""

import io
import os
import logging
from typing import Tuple

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document

logger = logging.getLogger(__name__)


def generate_pdf_bytes(doc_title: str, doc_text: str) -> bytes:
    """
    Renders document text into a styled PDF in memory using ReportLab.
    Returns bytes.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        alignment=1,  # Center
        spaceAfter=15,
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=15,
        spaceAfter=10,
    )

    story = []

    # Title
    story.append(Paragraph(doc_title, title_style))
    story.append(Spacer(1, 10))

    # Split text by lines and wrap in Paragraphs
    lines = doc_text.split('\n')
    for line in lines:
        line_str = line.strip()
        if not line_str:
            story.append(Spacer(1, 8))
            continue

        # Escape XML entities for ReportLab
        clean_text = line_str.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        story.append(Paragraph(clean_text, body_style))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_docx_bytes(doc_title: str, doc_text: str) -> bytes:
    """
    Renders document text into a Microsoft Word (.docx) file in memory.
    Returns bytes.
    """
    doc = Document()
    
    # Title
    title_para = doc.add_heading(doc_title, level=1)
    title_para.alignment = 1  # Center

    # Body
    lines = doc_text.split('\n')
    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue
        doc.add_paragraph(line_str)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
